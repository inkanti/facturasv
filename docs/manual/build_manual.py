#!/usr/bin/env python3
"""Genera el manual DOCX a partir del Markdown de docs/manual."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image

BASE = Path(__file__).resolve().parent
SOURCE = BASE / "manual_usuario_facturasv.md"
TARGET = BASE / "manual_usuario_facturasv.docx"
NAVY = RGBColor(15, 45, 107)
BLUE = RGBColor(18, 60, 140)
MUTED = RGBColor(95, 111, 135)
LIGHT = "EAF2FF"

CONTENT = SOURCE.read_text(encoding="utf-8")
BODY = CONTENT.split("<!-- contenido -->", 1)[1].strip()
LINES = BODY.splitlines()

# python-docx no acepta todos los JPEG; se normaliza el logo a RGB base.
LOGO = BASE / "img" / "bissweb.jpg"
PRINT_LOGO = BASE / "img" / "bissweb-print.jpg"
if LOGO.is_file() and not PRINT_LOGO.exists():
    with Image.open(LOGO) as source_image:
        source_image.convert("RGB").save(PRINT_LOGO, "JPEG", quality=92)


def set_cell_shading(cell, hex_color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def add_page_number(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    placeholder = OxmlElement("w:t")
    placeholder.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, placeholder, end])


def add_inline(paragraph, text: str, size: float | None = None, color=None) -> None:
    """Agrega texto con negritas y código inline."""
    text = text.replace("\\|", "|")
    parts = re.split(r"(\*\*.+?\*\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        bold = part.startswith("**") and part.endswith("**") and len(part) > 4
        code = part.startswith("`") and part.endswith("`") and len(part) > 2
        value = part[2:-2] if bold else (part[1:-1] if code else part)
        run = paragraph.add_run(value)
        run.bold = bold
        if code:
            run.font.name = "Consolas"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
            run.font.size = Pt((size or 10.5) - 1)
            run.font.color.rgb = RGBColor(30, 41, 59)
        else:
            if size:
                run.font.size = Pt(size)
            if color:
                run.font.color.rgb = color


def callout(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.25)
    paragraph.paragraph_format.right_indent = Cm(0.25)
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), LIGHT)
    p_pr.append(shd)
    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "18")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "2F8BFF")
    borders.append(left)
    p_pr.append(borders)
    add_inline(paragraph, text)
    for run in paragraph.runs:
        run.font.size = Pt(9.8)


def add_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    width = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=width)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for row_index, row in enumerate(rows):
        for col_index in range(width):
            cell = table.cell(row_index, col_index)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            value = row[col_index] if col_index < len(row) else ""
            add_inline(paragraph, value.strip(), size=8.6)
            if row_index == 0:
                set_cell_shading(cell, "0F2D6B")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "F5F8FD")
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_image(doc: Document, path: Path, caption: str) -> None:
    if not path.is_file():
        add_inline(doc.add_paragraph(), f"[Imagen no disponible: {path.name}]", size=9, color=MUTED)
        return
    with Image.open(path) as image:
        width, height = image.size
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    target_width = Cm(16.0)
    if height > width:
        target_width = Cm(9.5)
    run.add_picture(str(path), width=target_width)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_run = caption_paragraph.add_run(caption)
    caption_run.italic = True
    caption_run.font.size = Pt(8.2)
    caption_run.font.color.rgb = MUTED
    caption_paragraph.paragraph_format.space_after = Pt(10)


def parse_table(start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start
    while index < len(LINES) and LINES[index].strip().startswith("|"):
        raw = LINES[index].strip().strip("|")
        cells = [cell.strip() for cell in raw.split("|")]
        if not all(re.fullmatch(r":?-{2,}:?", cell) for cell in cells):
            rows.append(cells)
        index += 1
    return rows, index


def build() -> None:
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.different_first_page_header_footer = True
    section.first_page_header.paragraphs[0].text = ""
    section.first_page_footer.paragraphs[0].text = ""

    settings = doc.settings.element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(30, 41, 59)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)

    for name, size, color, before, after in (
        ("Heading 1", 19, NAVY, 16, 8),
        ("Heading 2", 13.5, BLUE, 12, 5),
        ("Heading 3", 11.5, RGBColor(30, 58, 138), 9, 4),
    ):
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    # Encabezado y pie
    header = section.header.paragraphs[0]
    header.text = "FacturaSV · Manual de usuario"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer.add_run("BissWeb · Impulso Digital   ·   Página ")
    footer_run.font.size = Pt(8)
    footer_run.font.color.rgb = MUTED
    add_page_number(footer)
    for run in footer.runs:
        run.font.size = Pt(8)
        run.font.color.rgb = MUTED

    # Portada
    for _ in range(2):
        doc.add_paragraph()
    logo = doc.add_paragraph()
    logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    logo.add_run().add_picture(str(PRINT_LOGO), width=Cm(4.4))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    run = title.add_run("Manual de usuario")
    run.font.size = Pt(32)
    run.font.bold = True
    run.font.color.rgb = NAVY
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = subtitle.add_run("FacturaSV")
    run.font.size = Pt(22)
    run.font.bold = True
    run.font.color.rgb = BLUE
    description = doc.add_paragraph()
    description.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = description.add_run("Sistema de facturación electrónica multiempresa para El Salvador")
    run.font.size = Pt(12)
    run.font.color.rgb = MUTED
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = meta.add_run("Módulos 1 al 14  ·  Esquema 15  ·  Octubre de 2026")
    run.font.size = Pt(10)
    run.font.color.rgb = MUTED
    brand = doc.add_paragraph()
    brand.alignment = WD_ALIGN_PARAGRAPH.CENTER
    brand.paragraph_format.space_before = Pt(28)
    run = brand.add_run("Desarrollado por BissWeb — Impulso Digital")
    run.font.size = Pt(9.5)
    run.font.bold = True
    run.font.color.rgb = NAVY
    doc.add_page_break()

    # Control de versión
    h = doc.add_heading("Control de versión", level=2)
    add_table(doc, [
        ["Versión", "Fecha", "Cambio"],
        ["1.0", "Octubre de 2026", "Manual de usuario de los módulos 1 al 14"],
        ["1.1", "Octubre de 2026", "Acceso con usuario o correo, sesión persistente y avisos de inventario"],
    ])

    # Índice generado desde los capítulos
    doc.add_heading("Contenido", level=1)
    chapters = [
        line[2:].strip() for line in LINES
        if line.startswith("# ") and not line.startswith("## ")
    ]
    for chapter in chapters:
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.space_after = Pt(3)
        run = paragraph.add_run(chapter)
        run.font.size = Pt(10.5)
        run.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_page_break()

    # Cuerpo, omitiendo el bloque de contenido del Markdown
    in_toc = False
    index = 0
    first_heading = True
    while index < len(LINES):
        raw = LINES[index]
        line = raw.strip()
        if not line:
            index += 1
            continue
        if line == "<!-- pagebreak -->":
            doc.add_page_break()
            index += 1
            continue
        if line.startswith("# "):
            title_text = line[2:].strip()
            if title_text.startswith("1. Introducción"):
                in_toc = False
            if in_toc:
                index += 1
                continue
            heading = doc.add_heading(title_text, level=1)
            if not first_heading:
                heading.paragraph_format.page_break_before = True
            first_heading = False
            index += 1
            continue
        if line.startswith("## "):
            if line[3:].strip() == "Contenido":
                in_toc = True
                index += 1
                continue
            if in_toc:
                index += 1
                continue
            doc.add_heading(line[3:].strip(), level=2)
            index += 1
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            index += 1
            continue
        if line.startswith("!["):
            match = re.match(r"!\[(.*?)\]\((.+?)\)", line)
            if match:
                add_image(doc, BASE / match.group(2), match.group(1))
            index += 1
            continue
        if line.startswith("|"):
            rows, next_index = parse_table(index)
            add_table(doc, rows)
            index = next_index
            continue
        if line.startswith(">"):
            callout(doc, line[1:].strip())
            index += 1
            continue
        if line.startswith("- "):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(3)
            add_inline(paragraph, line[2:].strip())
            index += 1
            continue
        numbered = re.match(r"^(\d+)\.\s+(.*)$", line)
        if numbered:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.left_indent = Cm(0.65)
            paragraph.paragraph_format.first_line_indent = Cm(-0.45)
            paragraph.paragraph_format.space_after = Pt(3)
            run = paragraph.add_run(numbered.group(1) + ". ")
            run.bold = True
            add_inline(paragraph, numbered.group(2))
            index += 1
            continue

        paragraph = doc.add_paragraph()
        add_inline(paragraph, line)
        index += 1

    # Propiedades
    doc.core_properties.title = "Manual de usuario FacturaSV"
    doc.core_properties.subject = "Sistema de facturación electrónica multiempresa"
    doc.core_properties.author = "BissWeb — Impulso Digital"
    doc.core_properties.language = "es-SV"
    doc.core_properties.comments = "Módulos 1 al 14; esquema 15"

    doc.save(TARGET)
    print(f"OK {TARGET}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)} images={len(doc.inline_shapes)}")


if __name__ == "__main__":
    build()
